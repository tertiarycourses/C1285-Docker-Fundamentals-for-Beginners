"""Post-convert fixes for the C1285 deck (run after convert_deck.py): one-day structure."""
from pptx import Presentation

DECK = "courseware/Docker-Fundamentals-for-Beginners-v1.pptx"
prs = Presentation(DECK)
S = list(prs.slides)


def set_para(para, text):
    para.runs[0].text = text
    for r in para.runs[1:]:
        r.text = ""


def shape(slide, sid):
    return next(s for s in slide.shapes if s.shape_id == sid)


def set_shape(slide, sid, lines):
    paras = shape(slide, sid).text_frame.paragraphs
    assert len(paras) == len(lines), (sid, len(paras), lines)
    for p, t in zip(paras, lines):
        set_para(p, t)


def replace_in(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if old in t:
                set_para(p, t.replace(old, new))
                n += 1
    assert n, (old, slide.slide_id)


# Cover
replace_in(S[0], "WSQ Course Code: C1285", "Course Code: C1285")

# Schedule slide -> one day, Morning / Afternoon
assert shape(S[6], 5).text_frame.text == "Lesson Plan — 2 Days, 8 hours/day"
set_shape(S[6], 5, ["Lesson Plan — 1 Day, 7.5 hours"])
set_shape(S[6], 9, ["Morning"])
set_shape(S[6], 10, ["Docker",
                     "· Fundamentals & commands (Labs 1–2)",
                     "· Dockerfile, build & CMD/ENTRYPOINT (Labs 3–5)",
                     "· Storage, networking, config (Labs 6–8)",
                     "· Sharing images on Docker Hub (Lab 9)"])
set_shape(S[6], 13, ["Afternoon"])
set_shape(S[6], 14, ["Docker Compose & Kubernetes",
                     "· Docker Compose (Labs 10–12)",
                     "· Pods & Namespaces (Labs 13–14)",
                     "· Deployments & Rollouts (Labs 15–16)",
                     "· Services (Lab 17)",
                     "· Storage, Jobs & CronJobs (Labs 18–19)"])
set_shape(S[6], 16, ["Timing:   9:30am – 5:30pm   ·   30-minute lunch (12:55pm – 1:25pm)"])

# Part dividers + lab eyebrows (no "Day 2" in a one-day course)
for s in S:
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if t in ("DAY 1", "DAY 2"):
                set_para(p, t.replace("DAY", "PART"))
            elif t.startswith(("DAY 1 · ", "DAY 2 · ")):
                set_para(p, t.replace("DAY ", "PART ", 1))

# "Day 2" forward references in lab notes
replace_in(S[54], " (and in Kubernetes, with ConfigMaps and Secrets — Day 2).", " (and in Kubernetes, with ConfigMaps and Secrets).")
replace_in(S[58], "how Kubernetes gets your app on Day 2:", "how Kubernetes gets your app this afternoon:")
replace_in(S[71], "deploy to Kubernetes on Day 2.", "deploy to Kubernetes later today.")
replace_in(S[74], "On Day 2 you'll deploy", "Later today you'll deploy")

# Closing divider
set_shape(S[121], 6, ["Wrap-up & Feedback"])

# One 30-min lunch, moved to after Lab 9 (before the Docker Compose topic divider)
lunch = S[37]
assert "Lunch Break" in [sh.text_frame.text for sh in lunch.shapes if sh.has_text_frame]
set_shape(lunch, 7, ["30 minutes"])
assert "TOPIC 6" in [sh.text_frame.text for sh in S[60].shapes if sh.has_text_frame]
lst = prs.slides._sldIdLst
ids = list(lst)
el = ids[37]
lst.remove(el)
ids[60].addprevious(el)

# Renumber the static bottom-right page-number footers.
W, H = prs.slide_width, prs.slide_height
for n, s in enumerate(prs.slides, start=1):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit() \
           and sh.top > H * 0.85 and sh.left > W * 0.5:
            set_para(sh.text_frame.paragraphs[0], str(n))
prs.save(DECK)
print("slides:", len(prs.slides))
# (applied after the first pass) slide 50 body: "...by name:tag (Day 2)." -> drop "(Day 2)"
