"""In-place WSQ -> non-WSQ conversion of the LP and LG for C1285.

Source: source-wsq/courseware/{LP,LG}-Application-Integration-with-Docker-and-Kubernetes.docx
Output: courseware/{LP,LG}-Docker-Fundamentals-for-Beginners.docx

The WSQ course is 2 days; C1285 is ONE day (9:30am-5:30pm, 450 min + one 30-min lunch).
Every WSQ topic/lab row is kept and rescaled proportionally (build/schedule.py).
"""
import copy
import sys

from docx import Document

sys.path.insert(0, "build")
from schedule import ROWS, cut  # noqa: E402

OLD_TITLE = "Application Integration with Docker and Kubernetes"
NEW_TITLE = "Docker Fundamentals for Beginners"
OLD_REPO = "https://github.com/tertiarycourses/TGS-2021010366-Application-Integration-with-Docker-and-Kubernetes"
NEW_REPO = "https://github.com/tertiarycourses/C1285-Docker-Fundamentals-for-Beginners"
SRC = "source-wsq/courseware/{}-Application-Integration-with-Docker-and-Kubernetes.docx"
OUT = "courseware/{}-Docker-Fundamentals-for-Beginners.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def drop(el):
    el.getparent().remove(el)


def set_text(container, text):
    paras = container.paragraphs if hasattr(container, "paragraphs") else [container]
    p = paras[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)
    for extra in paras[1:]:
        drop(extra._element)


def sub(p, old, new):
    """Replace inside a single run when possible (keeps bold labels); else collapse onto run 0."""
    if old not in p.text:
        return False
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    t = p.text.replace(old, new)
    p.runs[0].text = t
    for r in p.runs[1:]:
        r.text = ""
    print("  collapsed runs:", t[:80])
    return True


def all_paras(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


def common(doc, version_summary):
    for p in all_paras(doc):
        sub(p, "TGS Ref No: TGS-2021010366", "Course Code: C1285")
        sub(p, "TGS-2021010366", "C1285")
        sub(p, OLD_REPO, NEW_REPO)
        sub(p, OLD_TITLE, NEW_TITLE)
    rels = doc.part.rels
    for h in doc.element.body.iter(W + "hyperlink"):
        rid = h.get(R + "id")
        if rid and rels[rid].is_external and OLD_REPO in rels[rid].target_ref:
            rels[rid]._target = rels[rid].target_ref.replace(OLD_REPO, NEW_REPO)
            for t in h.iter(W + "t"):
                t.text = t.text.replace(OLD_REPO, NEW_REPO)
    vt = doc.tables[0]
    assert vt.rows[0].cells[0].text == "Version Number"
    for c, v in zip(vt.rows[1].cells, ["1.0", "3 October 2026", version_summary, "Tertiary Infotech Academy Pte Ltd"]):
        set_text(c, v)
    for r in list(vt.rows)[2:]:
        drop(r._tr)


def delete_section(doc, heading):
    start = next(p for p in doc.paragraphs if p.text == heading and p.style.name == "Heading 1")
    el = start._element
    nxt = el.getnext()
    drop(el)
    while nxt is not None:
        if nxt.tag == W + "p":
            sty = nxt.find(f".//{W}pStyle")
            if sty is not None and sty.get(W + "val") in ("Heading1",):
                break
        if nxt.tag == W + "sectPr":
            break
        cur, nxt = nxt, nxt.getnext()
        drop(cur)


# ------------------------------------------------------------------ LG ------
lg = Document(SRC.format("LG"))
common(lg, "First version — step-by-step guide to all 19 Docker & Kubernetes labs for the one-day course.")
LG_SUBS = [
    ("in the WSQ course ", "in the course "),
    ("Over two days you build", "Over one day you build"),
    ("Course flow at a glance — ", "Course flow at a glance — "),
    ("Day 1 (Docker):", "Morning (Docker):"),
    ("Docker Hub, Docker Compose (Labs 1–12).", "Docker Hub (Labs 1–9)."),
    ("Day 2 (Kubernetes):", "Afternoon (Docker Compose & Kubernetes): Docker Compose (Labs 10–12),"),
    ("(all Day 1 labs)", "(all Docker labs)"),
    ("All Day 2 Kubernetes labs", "All Kubernetes labs"),
    ("built-in Kubernetes) for Day 2.", "built-in Kubernetes) for the Kubernetes labs."),
    (" (and in Kubernetes, with ConfigMaps and Secrets — Day 2).", " (and in Kubernetes, with ConfigMaps and Secrets)."),
    ("how Kubernetes gets your app on Day 2:", "how Kubernetes gets your app this afternoon:"),
    ("deploy to Kubernetes on Day 2.", "deploy to Kubernetes later today."),
    ("On Day 2 you'll deploy", "Later today you'll deploy"),
]
for p in all_paras(lg):
    for a, b in LG_SUBS:
        sub(p, a, b)
lg.save(OUT.format("LG"))

# ------------------------------------------------------------------ LP ------
lp = Document(SRC.format("LP"))
common(lp, "First version — one-day lesson plan (Docker + Kubernetes), 7.5 instructional hours.")
for p in lp.paragraphs:
    sub(p, "This 2-day hands-on course", "This one-day hands-on course")
    sub(p, "Day 1 covers Docker: commands, building images with a Dockerfile, CMD vs ENTRYPOINT, volumes, "
           "networking, configuration, Docker Hub and Docker Compose. Day 2 covers Kubernetes:",
           "The morning covers Docker: commands, building images with a Dockerfile, CMD vs ENTRYPOINT, volumes, "
           "networking, configuration and Docker Hub. The afternoon covers Docker Compose and Kubernetes:")
    if p.text == "Docker — Images, Storage, Networking & Compose":
        set_text(p, "Docker & Kubernetes — Build, Compose, Deploy & Scale")

facts = lp.tables[1]
assert facts.rows[0].cells[0].text == "Course Title"
for r in list(facts.rows):
    key = r.cells[0].text
    if key == "Course Code":
        set_text(r.cells[1], "C1285")
    elif key == "Duration":
        set_text(r.cells[1], "1 Day (7.5 instructional hours)")
    elif key == "Daily Schedule":
        set_text(r.cells[1], "9:30 AM – 5:30 PM (7.5 instructional hours, excluding lunch)")
    elif key == "Lunch Break":
        set_text(r.cells[1], "12:55 PM – 1:25 PM (30 minutes)")
    elif key == "Breaks":
        drop(r._tr)

# Day 2 block: heading, subtitle, two session lines and its three tables.
paras = lp.paragraphs
d2 = next(i for i, p in enumerate(paras) if p.text == "Day 2" and p.style.name == "Heading 1")
d2_el = paras[d2]._element
labs = next(p for p in paras if p.text == "Labs Covered")._element
el = d2_el
while el is not labs:
    nxt = el.getnext()
    drop(el)
    el = nxt
delete_section(lp, "Assessment")

d1 = next(p for p in lp.paragraphs if p.text == "Day 1" and p.style.name == "Heading 1")
set_text(d1, "Day 1 · 9:30am – 5:30pm")
for p in lp.paragraphs:
    if p.text.startswith("Morning Session"):
        set_text(p, "Morning Session  ·  9:30 AM – 12:55 PM")
    elif p.text.startswith("Afternoon Session"):
        set_text(p, "Afternoon Session  ·  1:25 PM – 5:30 PM")

morning, lunch, afternoon = lp.tables[2], lp.tables[3], lp.tables[4]
assert morning.rows[1].cells[1].text.startswith("Welcome")
assert lunch.rows[1].cells[1].text == "Lunch Break"
assert afternoon.rows[1].cells[1].text.startswith("Topic 3")
body = copy.deepcopy(morning.rows[2]._tr)      # an ordinary teaching row (not the tea break)


def clock(t):
    h, m = divmod(t, 60)
    return f"{(h - 1) % 12 + 1}:{m:02d}"


def fill(table, rows, start):
    for r in list(table.rows)[1:]:
        drop(r._tr)
    t = start
    for label, mins in rows:
        table._tbl.append(copy.deepcopy(body))
        cells = table.rows[-1].cells
        for c, v in zip(cells, [f"{clock(t)} – {clock(t + mins)}", label, f"{mins} min"]):
            set_text(c, v)
        t += mins
    return t


end_am = fill(morning, ROWS[:cut], 9 * 60 + 30)
assert end_am == 12 * 60 + 55, end_am
set_text(lunch.rows[1].cells[0], "12:55 – 1:25")
set_text(lunch.rows[1].cells[2], "30 min")
end_pm = fill(afternoon, ROWS[cut:], 13 * 60 + 25)
assert end_pm == 17 * 60 + 30, end_pm
lp.save(OUT.format("LP"))
print("ok")
